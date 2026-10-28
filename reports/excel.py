import io

from openpyxl import Workbook
from openpyxl.styles import Font


def render_client_report_excel(report):
    """Renders the dict from reports.services.build_client_report() as an .xlsx (bytes)."""
    wb = Workbook()
    _write_summary_sheet(wb.active, report)
    _write_rules_sheet(wb.create_sheet("Regras"), report["rules"])
    _write_gaps_sheet(wb.create_sheet("Gaps"), report["gaps"])
    _write_usecases_sheet(wb.create_sheet("Casos de Uso"), report["use_case_links"])
    _write_logsources_sheet(wb.create_sheet("Fontes de Log"), report["log_source_links"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _set_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _header_row(ws, labels):
    for col, label in enumerate(labels, start=1):
        ws.cell(row=1, column=col, value=label).font = Font(bold=True)


def _date(value, fmt="%d/%m/%Y"):
    return value.strftime(fmt) if value else ""


def _write_summary_sheet(ws, report):
    ws.title = "Resumo"
    client = report["client"]
    coverage = report["coverage"]
    rows = [
        ("Cliente", client.name),
        ("Gerado em", report["generated_at"].strftime("%d/%m/%Y %H:%M")),
        ("Regras cadastradas", len(report["rules"])),
        ("Regras ativas", report["active_rules_count"]),
        ("Gaps em aberto", report["open_gaps_count"]),
        ("Gaps cadastrados", len(report["gaps"])),
        (
            "Casos de uso cobertos",
            f"{coverage['covered']} de {coverage['applicable']}" if coverage["applicable"] else "-",
        ),
        ("Percentual de cobertura", f"{coverage['percentage']}%" if coverage["percentage"] is not None else "-"),
    ]
    for row, (label, value) in enumerate(rows, start=1):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
    _set_widths(ws, {"A": 28, "B": 40})


def _write_rules_sheet(ws, rules):
    _header_row(ws, ["ID", "Nome", "Tecnologia", "Severidade", "Status", "Fonte de Log", "Última revisão", "Próxima revisão"])
    for row, rule in enumerate(rules, start=2):
        ws.cell(row=row, column=1, value=rule.internal_id)
        ws.cell(row=row, column=2, value=rule.name)
        ws.cell(row=row, column=3, value=str(rule.technology))
        ws.cell(row=row, column=4, value=str(rule.severity) if rule.severity else "")
        ws.cell(row=row, column=5, value=str(rule.status) if rule.status else "")
        ws.cell(row=row, column=6, value=str(rule.log_source) if rule.log_source else "")
        ws.cell(row=row, column=7, value=_date(rule.last_reviewed_at))
        ws.cell(row=row, column=8, value=_date(rule.next_review_at))
    _set_widths(ws, {"A": 16, "B": 38, "C": 16, "D": 14, "E": 16, "F": 22, "G": 14, "H": 14})


def _write_gaps_sheet(ws, gaps):
    _header_row(ws, ["Código", "Título", "Status", "Criticidade", "Responsável", "Aberto em", "Concluído em"])
    for row, gap in enumerate(gaps, start=2):
        ws.cell(row=row, column=1, value=gap.code)
        ws.cell(row=row, column=2, value=gap.title)
        ws.cell(row=row, column=3, value=str(gap.status) if gap.status else "")
        ws.cell(row=row, column=4, value=str(gap.criticality) if gap.criticality else "")
        ws.cell(row=row, column=5, value=str(gap.responsible_party) if gap.responsible_party else "")
        ws.cell(row=row, column=6, value=_date(gap.opened_at))
        ws.cell(row=row, column=7, value=_date(gap.closed_at))
    _set_widths(ws, {"A": 14, "B": 38, "C": 16, "D": 14, "E": 16, "F": 14, "G": 14})


def _write_usecases_sheet(ws, use_case_links):
    _header_row(ws, ["Caso de Uso", "Status", "Criticidade", "Responsável", "Última validação"])
    for row, link in enumerate(use_case_links, start=2):
        ws.cell(row=row, column=1, value=str(link.use_case))
        ws.cell(row=row, column=2, value=str(link.status) if link.status else "")
        ws.cell(row=row, column=3, value=str(link.criticality) if link.criticality else "")
        ws.cell(row=row, column=4, value=str(link.owner) if link.owner else "")
        ws.cell(row=row, column=5, value=_date(link.last_validated_at))
    _set_widths(ws, {"A": 32, "B": 16, "C": 14, "D": 18, "E": 16})


def _write_logsources_sheet(ws, log_source_links):
    _header_row(ws, ["Fonte de Log", "Status", "Último evento"])
    for row, link in enumerate(log_source_links, start=2):
        ws.cell(row=row, column=1, value=str(link.log_source))
        ws.cell(row=row, column=2, value=str(link.status) if link.status else "")
        ws.cell(row=row, column=3, value=_date(link.last_event_at, "%d/%m/%Y %H:%M"))
    _set_widths(ws, {"A": 32, "B": 16, "C": 18})
