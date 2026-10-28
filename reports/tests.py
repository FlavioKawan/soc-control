import io

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from catalog.models import ClientLogSource, ClientUseCase, LogSource, UseCase
from clients.models import Client
from detections.models import Rule
from gaps.models import Gap
from taxonomy.choices import StatusCategory, StatusScope
from taxonomy.models import Framework, FrameworkControl, ResponsibleParty, Status, Technology

from .services import FRAMEWORK_FULL_LIST_MAX_CONTROLS, build_client_facing_report, build_client_report


class ClientReportServiceTests(TestCase):
    def setUp(self):
        self.client_obj = Client.objects.create(name="Acme Corp")
        technology = Technology.objects.create(name="Wazuh")
        active = Status.objects.create(name="Ativa", scope=[StatusScope.RULE], category=StatusCategory.ACTIVE)
        Rule.objects.create(
            internal_id="SOC-001", name="Brute Force", client=self.client_obj, technology=technology, status=active
        )
        Rule.objects.create(internal_id="SOC-002", name="Port Scan", client=self.client_obj, technology=technology)
        Gap.objects.create(code="GAP-001", client=self.client_obj, title="DNS logs ausentes")

    def test_report_counts_match_the_clients_data(self):
        report = build_client_report(self.client_obj)
        self.assertEqual(len(report["rules"]), 2)
        self.assertEqual(report["active_rules_count"], 1)
        self.assertEqual(len(report["gaps"]), 1)
        self.assertEqual(report["open_gaps_count"], 1)

    def test_coverage_is_none_without_applicable_use_cases(self):
        report = build_client_report(self.client_obj)
        self.assertIsNone(report["coverage"]["percentage"])

    def test_coverage_reflects_covered_use_cases(self):
        covered = Status.objects.create(
            name="Coberto", scope=[StatusScope.CLIENT_USE_CASE], category=StatusCategory.COVERED
        )
        use_case = UseCase.objects.create(code="UC-001", name="Brute Force")
        ClientUseCase.objects.create(client=self.client_obj, use_case=use_case, status=covered)

        report = build_client_report(self.client_obj)
        self.assertEqual(report["coverage"]["percentage"], 100)


class ClientFacingReportServiceTests(TestCase):
    def setUp(self):
        self.client_obj = Client.objects.create(name="Acme Corp")
        self.technology = Technology.objects.create(name="Wazuh")
        self.active = Status.objects.create(name="Ativa", scope=[StatusScope.RULE], category=StatusCategory.ACTIVE)
        self.testing = Status.objects.create(name="Teste", scope=[StatusScope.RULE], category=StatusCategory.TESTING)

    def _rule(self, internal_id, status, **kwargs):
        return Rule.objects.create(
            internal_id=internal_id, name=internal_id, client=self.client_obj,
            technology=self.technology, status=status, **kwargs,
        )

    def test_only_active_rules_are_listed(self):
        self._rule("SOC-001", self.active)
        self._rule("SOC-002", self.testing)

        report = build_client_facing_report(self.client_obj)
        self.assertEqual([rule.internal_id for rule in report["active_rules"]], ["SOC-001"])

    def test_gaps_are_split_by_whether_the_client_is_responsible(self):
        cliente = ResponsibleParty.objects.create(name="Cliente", is_client=True)
        soc = ResponsibleParty.objects.create(name="SOC")
        closed = Status.objects.create(name="Concluído", scope=[StatusScope.GAP], category=StatusCategory.CLOSED)
        Gap.objects.create(code="GAP-1", client=self.client_obj, title="A", responsible_party=cliente)
        Gap.objects.create(code="GAP-2", client=self.client_obj, title="B", responsible_party=soc)
        Gap.objects.create(code="GAP-3", client=self.client_obj, title="C")
        Gap.objects.create(code="GAP-4", client=self.client_obj, title="D", responsible_party=cliente, status=closed)

        report = build_client_facing_report(self.client_obj)
        self.assertEqual([gap.code for gap in report["client_gaps"]], ["GAP-1"])
        self.assertEqual(sorted(gap.code for gap in report["other_gaps"]), ["GAP-2", "GAP-3"])
        self.assertEqual(report["open_gaps_count"], 3)

    def test_problem_log_sources_report_how_many_active_rules_they_affect(self):
        not_receiving = Status.objects.create(
            name="Sem log", scope=[StatusScope.CLIENT_LOG_SOURCE], category=StatusCategory.NOT_RECEIVING
        )
        receiving = Status.objects.create(
            name="Recebendo", scope=[StatusScope.CLIENT_LOG_SOURCE], category=StatusCategory.RECEIVING
        )
        broken = LogSource.objects.create(name="Firewall")
        healthy = LogSource.objects.create(name="AD")
        ClientLogSource.objects.create(client=self.client_obj, log_source=broken, status=not_receiving)
        ClientLogSource.objects.create(client=self.client_obj, log_source=healthy, status=receiving)
        self._rule("SOC-001", self.active, log_source=broken)
        self._rule("SOC-002", self.testing, log_source=broken)

        report = build_client_facing_report(self.client_obj)
        self.assertEqual([link.log_source for link in report["problem_log_sources"]], [broken])
        self.assertEqual(report["problem_log_sources"][0].affected_rules_count, 1)
        self.assertEqual((report["log_sources_receiving"], report["log_sources_total"]), (1, 2))

    def test_framework_coverage_counts_only_this_clients_active_rules(self):
        framework = Framework.objects.create(name="Cyber Kill Chain")
        recon = FrameworkControl.objects.create(framework=framework, code="KC1", name="Reconhecimento")
        FrameworkControl.objects.create(framework=framework, code="KC2", name="Armamento")
        self._rule("SOC-001", self.active).framework_controls.add(recon)
        self._rule("SOC-002", self.testing).framework_controls.add(recon)
        other_client = Client.objects.create(name="Outro")
        Rule.objects.create(
            internal_id="X-1", name="X", client=other_client, technology=self.technology, status=self.active
        ).framework_controls.add(recon)

        [item] = build_client_facing_report(self.client_obj)["frameworks"]
        self.assertEqual((item["covered_count"], item["total"]), (1, 2))
        self.assertTrue(item["lists_all"])
        self.assertEqual([(c.code, c.active_rules_count) for c in item["controls"]], [("KC1", 1), ("KC2", 0)])

    def test_large_frameworks_list_only_covered_controls(self):
        framework = Framework.objects.create(name="MITRE ATT&CK")
        controls = FrameworkControl.objects.bulk_create(
            FrameworkControl(framework=framework, code=f"T{i:04d}", name=f"Técnica {i}")
            for i in range(FRAMEWORK_FULL_LIST_MAX_CONTROLS + 1)
        )
        self._rule("SOC-001", self.active).framework_controls.add(controls[0])

        [item] = build_client_facing_report(self.client_obj)["frameworks"]
        self.assertFalse(item["lists_all"])
        self.assertEqual([c.code for c in item["controls"]], ["T0000"])

    def test_not_applicable_use_cases_are_left_out(self):
        not_applicable = Status.objects.create(
            name="N/A", scope=[StatusScope.CLIENT_USE_CASE], category=StatusCategory.NOT_APPLICABLE
        )
        covered = Status.objects.create(
            name="Coberto", scope=[StatusScope.CLIENT_USE_CASE], category=StatusCategory.COVERED
        )
        ClientUseCase.objects.create(
            client=self.client_obj, use_case=UseCase.objects.create(code="UC-1", name="A"), status=covered
        )
        ClientUseCase.objects.create(
            client=self.client_obj, use_case=UseCase.objects.create(code="UC-2", name="B"), status=not_applicable
        )

        report = build_client_facing_report(self.client_obj)
        self.assertEqual([link.use_case.code for link in report["use_case_links"]], ["UC-1"])


class ClientReportViewTests(TestCase):
    def setUp(self):
        call_command("setup_roles")
        self.client_obj = Client.objects.create(name="Acme Corp")
        technology = Technology.objects.create(name="Wazuh")
        Rule.objects.create(internal_id="SOC-001", name="Brute Force", client=self.client_obj, technology=technology)
        Gap.objects.create(code="GAP-001", client=self.client_obj, title="DNS logs ausentes")

    def test_pdf_export_requires_login(self):
        response = self.client.get(reverse("reports:client_report_pdf", kwargs={"pk": self.client_obj.pk}))
        self.assertEqual(response.status_code, 302)

    def test_viewer_can_download_pdf(self):
        user = User.objects.create_user("viewer", password="x")
        user.groups.add(Group.objects.get(name="VIEWER"))
        self.client.force_login(user)

        response = self.client.get(reverse("reports:client_report_pdf", kwargs={"pk": self.client_obj.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_viewer_can_download_excel_and_it_contains_expected_sheets(self):
        user = User.objects.create_user("viewer2", password="x")
        user.groups.add(Group.objects.get(name="VIEWER"))
        self.client.force_login(user)

        response = self.client.get(reverse("reports:client_report_excel", kwargs={"pk": self.client_obj.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Resumo", "Regras", "Gaps", "Casos de Uso", "Fontes de Log"])
        self.assertEqual(wb["Regras"].cell(row=2, column=1).value, "SOC-001")
        self.assertEqual(wb["Gaps"].cell(row=2, column=1).value, "GAP-001")

    def test_viewer_can_download_client_facing_pdf(self):
        user = User.objects.create_user("viewer3", password="x")
        user.groups.add(Group.objects.get(name="VIEWER"))
        self.client.force_login(user)
        Gap.objects.create(code="GAP-002", client=self.client_obj, title="Título com <marcação> & símbolos")

        response = self.client.get(
            reverse("reports:client_facing_report_pdf", kwargs={"pk": self.client_obj.pk})
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn("relatorio-cliente-acme-corp.pdf", response["Content-Disposition"])

    def test_client_facing_pdf_requires_view_client_permission(self):
        user = User.objects.create_user("nobody2", password="x")
        self.client.force_login(user)

        response = self.client.get(
            reverse("reports:client_facing_report_pdf", kwargs={"pk": self.client_obj.pk})
        )
        self.assertEqual(response.status_code, 403)

    def test_user_without_view_client_permission_gets_403(self):
        user = User.objects.create_user("nobody", password="x")
        self.client.force_login(user)

        response = self.client.get(reverse("reports:client_report_pdf", kwargs={"pk": self.client_obj.pk}))
        self.assertEqual(response.status_code, 403)
